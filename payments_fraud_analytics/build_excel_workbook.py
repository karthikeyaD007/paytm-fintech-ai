"""
Builds merchant_workbook.xlsx from merchants.csv / ledger.csv, matching the
capstone grading rubric precisely.

Sheets produced:
  Merchants              - raw merchant reference data (VLOOKUP source, A2:D41)
  Fee_Tiers              - horizontal payment-method fee table (HLOOKUP source)
  Transactions_View      - one row per transaction (547 rows) with live VLOOKUP /
                            HLOOKUP formula columns; also a real Excel Table (ListObject)
  Merchant_Day_Analysis  - pivot-style merchant x day aggregation (static values,
                            computed with pandas) + nested IF/AND classification formula
  Merchant_Status_Pivot  - pivot-style merchant x status tables (sum + count) and a
                            unique-days-transacted vs total-transaction-count comparison

Rerun any time with `python build_excel_workbook.py` - it always rebuilds the
.xlsx from scratch.
"""

import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_PATH = os.path.join(BASE_DIR, "merchant_workbook.xlsx")

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
TITLE_FONT = Font(bold=True, size=12, color="305496")
NOTE_FONT = Font(bold=True, italic=True, color="C00000")

# Fee-tier percentages chosen for the HLOOKUP demo (MDR-style, illustrative).
FEE_TIERS = {"UPI": 0.000, "Wallet": 0.015, "Card": 0.018, "Netbanking": 0.012}
PAYMENT_METHODS = ["UPI", "Wallet", "Card", "Netbanking"]

# Nested IF/AND rule for "High-Value Merchant Day".
HV_THRESHOLD = 5000
HV_EXCLUDED_REGION = "East"


def style_header_row(ws, row, n_cols, col_offset=1):
    for c in range(col_offset, col_offset + n_cols):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def autosize(ws, n_cols, width=16, col_offset=1):
    for c in range(col_offset, col_offset + n_cols):
        ws.column_dimensions[get_column_letter(c)].width = width


def load_data():
    merchants = pd.read_csv(os.path.join(BASE_DIR, "merchants.csv"))
    ledger = pd.read_csv(os.path.join(BASE_DIR, "ledger.csv"))
    ledger["transaction_time"] = pd.to_datetime(ledger["transaction_time"])
    ledger["txn_date"] = ledger["transaction_time"].dt.date
    return merchants, ledger


# ---------------------------------------------------------------------------
# Merchants (VLOOKUP source)
# ---------------------------------------------------------------------------
def write_merchants_sheet(wb, merchants):
    ws = wb.create_sheet("Merchants")
    cols = list(merchants.columns)
    ws.append(cols)
    for row in merchants.itertuples(index=False):
        ws.append(list(row))
    style_header_row(ws, 1, len(cols))
    autosize(ws, len(cols))
    return ws


# ---------------------------------------------------------------------------
# Fee_Tiers (HLOOKUP source - horizontal layout)
# ---------------------------------------------------------------------------
def write_fee_tiers_sheet(wb):
    ws = wb.create_sheet("Fee_Tiers")
    ws["A1"] = (
        "Chosen MDR-style fee percentages (illustrative): "
        + ", ".join(f"{m}={FEE_TIERS[m]:.1%}" for m in PAYMENT_METHODS)
    )
    ws["A1"].font = NOTE_FONT
    ws.merge_cells("A1:F1")

    ws["A2"] = "Payment Method"
    ws["A3"] = "Fee % (MDR-style)"
    ws["A2"].font = Font(bold=True)
    ws["A3"].font = Font(bold=True)
    for i, method in enumerate(PAYMENT_METHODS):
        col = get_column_letter(2 + i)
        ws[f"{col}2"] = method
        fee_cell = ws[f"{col}3"]
        fee_cell.value = FEE_TIERS[method]
        fee_cell.number_format = "0.0%"
    style_header_row(ws, 2, len(PAYMENT_METHODS) + 1)
    autosize(ws, len(PAYMENT_METHODS) + 1, width=16)
    return ws


# ---------------------------------------------------------------------------
# Merchant_Day_Analysis (pivot-style merchant x day aggregation + nested IF/AND)
# ---------------------------------------------------------------------------
def build_merchant_day_df(merchants, ledger):
    daily = (
        ledger.groupby(["merchant_id", "txn_date"])["amount_inr"]
        .sum()
        .reset_index()
        .rename(columns={"amount_inr": "daily_total_inr"})
    )
    daily = daily.merge(
        merchants[["merchant_id", "merchant_name", "region"]], on="merchant_id", how="left"
    )
    daily["merchant_day_key"] = daily.apply(
        lambda r: f"{int(r['merchant_id'])}_{r['txn_date'].isoformat()}", axis=1
    )
    daily = daily.sort_values(["merchant_id", "txn_date"]).reset_index(drop=True)
    return daily


def write_merchant_day_analysis_sheet(wb, merchants, ledger):
    ws = wb.create_sheet("Merchant_Day_Analysis")
    ws["A1"] = (
        "PIVOT-STYLE AGGREGATION (static values, computed with pandas groupby by "
        "merchant_id x day — NOT a live Excel PivotTable, since openpyxl cannot "
        "create live PivotTable objects)."
    )
    ws["A1"].font = NOTE_FONT
    ws.merge_cells("A1:G1")

    ws["A2"] = (
        f'EXACT RULE: classification = "High-Value Merchant Day" when '
        f'daily_total_inr > {HV_THRESHOLD} AND region <> "{HV_EXCLUDED_REGION}"; otherwise "Normal".'
    )
    ws["A2"].font = NOTE_FONT
    ws.merge_cells("A2:G2")

    headers = [
        "merchant_day_key",
        "merchant_id",
        "merchant_name",
        "region",
        "txn_date",
        "daily_total_inr",
        "classification (nested IF/AND)",
    ]
    header_row = 4
    for c, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=c, value=h)
    style_header_row(ws, header_row, len(headers))

    daily = build_merchant_day_df(merchants, ledger)
    first_data_row = header_row + 1
    for i, row in daily.iterrows():
        r = first_data_row + i
        ws.cell(row=r, column=1, value=row["merchant_day_key"])
        ws.cell(row=r, column=2, value=int(row["merchant_id"]))
        ws.cell(row=r, column=3, value=row["merchant_name"])
        ws.cell(row=r, column=4, value=row["region"])
        date_cell = ws.cell(row=r, column=5, value=row["txn_date"])
        date_cell.number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=6, value=float(row["daily_total_inr"]))
        ws.cell(row=r, column=7, value=(
            f'=IF(AND(F{r}>{HV_THRESHOLD},D{r}<>"{HV_EXCLUDED_REGION}"),'
            f'"High-Value Merchant Day","Normal")'
        ))

    last_row = first_data_row + len(daily) - 1
    autosize(ws, len(headers), width=20)
    ws.freeze_panes = f"A{first_data_row}"
    return ws, last_row


# ---------------------------------------------------------------------------
# Transactions_View (VLOOKUP x3, HLOOKUP x1, link-back VLOOKUP, Excel Table)
# ---------------------------------------------------------------------------
def write_transactions_view_sheet(wb, merchants, ledger, day_analysis_last_row):
    ws = wb.create_sheet("Transactions_View")

    headers = [
        "transaction_id", "user_id", "merchant_id", "transaction_time",
        "amount_inr", "payment_method", "status", "risk_score",
        "merchant_name (VLOOKUP)", "category (VLOOKUP)", "region (VLOOKUP)",
        "fee_pct (HLOOKUP)", "fee_amount_inr", "txn_date", "merchant_day_key",
        "day_classification (VLOOKUP link-back)",
    ]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    n_merchants = len(merchants)
    merchants_last_row = n_merchants + 1  # header on row 1 of Merchants sheet

    ledger_sorted = ledger.sort_values("transaction_id").reset_index(drop=True)

    for i, row in ledger_sorted.iterrows():
        r = i + 2
        ws.cell(row=r, column=1, value=row["transaction_id"])
        ws.cell(row=r, column=2, value=int(row["user_id"]))
        ws.cell(row=r, column=3, value=int(row["merchant_id"]))
        dt_cell = ws.cell(row=r, column=4, value=row["transaction_time"].to_pydatetime())
        dt_cell.number_format = "yyyy-mm-dd hh:mm:ss"
        ws.cell(row=r, column=5, value=float(row["amount_inr"]))
        ws.cell(row=r, column=6, value=row["payment_method"])
        ws.cell(row=r, column=7, value=row["status"])
        ws.cell(row=r, column=8, value=int(row["risk_score"]))

        ws.cell(row=r, column=9, value=(
            f'=IFERROR(VLOOKUP(C{r},Merchants!$A$2:$D${merchants_last_row},2,FALSE),'
            f'"Merchant not found")'
        ))
        ws.cell(row=r, column=10, value=(
            f'=IFERROR(VLOOKUP(C{r},Merchants!$A$2:$D${merchants_last_row},3,FALSE),'
            f'"Merchant not found")'
        ))
        ws.cell(row=r, column=11, value=(
            f'=IFERROR(VLOOKUP(C{r},Merchants!$A$2:$D${merchants_last_row},4,FALSE),'
            f'"Merchant not found")'
        ))

        fee_cell = ws.cell(row=r, column=12, value=f'=HLOOKUP(F{r},Fee_Tiers!$B$2:$E$3,2,FALSE)')
        fee_cell.number_format = "0.0%"
        ws.cell(row=r, column=13, value=f'=E{r}*L{r}')

        date_cell = ws.cell(row=r, column=14, value=row["txn_date"])
        date_cell.number_format = "yyyy-mm-dd"

        ws.cell(row=r, column=15, value=f'=C{r}&"_"&TEXT(N{r},"yyyy-mm-dd")')

        ws.cell(row=r, column=16, value=(
            f'=IFERROR(VLOOKUP(O{r},Merchant_Day_Analysis!$A$5:$G${day_analysis_last_row},7,FALSE),'
            f'"Not found")'
        ))

    n_rows = len(ledger_sorted) + 1
    autosize(ws, len(headers), width=20)
    ws.freeze_panes = "A2"

    last_col = get_column_letter(len(headers))
    table = Table(displayName="TransactionsViewTable", ref=f"A1:{last_col}{n_rows}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    return ws


# ---------------------------------------------------------------------------
# Merchant_Status_Pivot (sum + count by merchant x status, and count-vs-unique)
# ---------------------------------------------------------------------------
def write_merchant_status_pivot_sheet(wb, merchants, ledger):
    ws = wb.create_sheet("Merchant_Status_Pivot")
    ws["A1"] = (
        "PIVOT-STYLE TABLES (static values, computed with pandas pivot_table/groupby — "
        "NOT live Excel PivotTables). See Transactions_View's Excel Table to build a "
        "live PivotTable in Excel yourself."
    )
    ws["A1"].font = NOTE_FONT
    ws.merge_cells("A1:H1")

    statuses = ["captured", "failed", "chargeback"]
    merged = ledger.merge(merchants[["merchant_id", "merchant_name"]], on="merchant_id", how="left")

    sum_pt = pd.pivot_table(
        merged, index=["merchant_id", "merchant_name"], columns="status",
        values="amount_inr", aggfunc="sum", fill_value=0,
    ).reindex(columns=statuses, fill_value=0).reset_index()

    count_pt = pd.pivot_table(
        merged, index=["merchant_id", "merchant_name"], columns="status",
        values="transaction_id", aggfunc="count", fill_value=0,
    ).reindex(columns=statuses, fill_value=0).reset_index()

    cursor = 3

    def write_table(title, df, cursor, value_fmt=None):
        ws.cell(row=cursor, column=1, value=title).font = TITLE_FONT
        cursor += 1
        header_row = cursor
        headers = ["merchant_id", "merchant_name"] + statuses
        for c, h in enumerate(headers, start=1):
            ws.cell(row=header_row, column=c, value=h)
        style_header_row(ws, header_row, len(headers))
        r = header_row + 1
        for _, row in df.iterrows():
            ws.cell(row=r, column=1, value=int(row["merchant_id"]))
            ws.cell(row=r, column=2, value=row["merchant_name"])
            for c, status in enumerate(statuses, start=3):
                cell = ws.cell(row=r, column=c, value=float(row[status]))
                if value_fmt:
                    cell.number_format = value_fmt
            r += 1
        return r + 2

    cursor = write_table("Pivot A: SUM(amount_inr) by merchant_id x status", sum_pt, cursor)
    cursor = write_table("Pivot B: COUNT(transactions) by merchant_id x status", count_pt, cursor)

    # Count-vs-count-unique comparison: unique days transacted vs total transaction count
    daily_counts = ledger.groupby("merchant_id").agg(
        unique_days_transacted=("txn_date", "nunique"),
        total_transaction_count=("transaction_id", "count"),
    ).reset_index()
    daily_counts = daily_counts.merge(
        merchants[["merchant_id", "merchant_name"]], on="merchant_id", how="left"
    ).sort_values("total_transaction_count", ascending=False)

    ws.cell(row=cursor, column=1, value=(
        "Comparison: unique days transacted vs. total transaction count (per merchant)"
    )).font = TITLE_FONT
    cursor += 1
    header_row = cursor
    headers = ["merchant_id", "merchant_name", "unique_days_transacted", "total_transaction_count"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=c, value=h)
    style_header_row(ws, header_row, len(headers))
    r = header_row + 1
    for _, row in daily_counts.iterrows():
        ws.cell(row=r, column=1, value=int(row["merchant_id"]))
        ws.cell(row=r, column=2, value=row["merchant_name"])
        ws.cell(row=r, column=3, value=int(row["unique_days_transacted"]))
        ws.cell(row=r, column=4, value=int(row["total_transaction_count"]))
        r += 1

    autosize(ws, 8, width=20)
    return ws


def main():
    if os.path.exists(OUT_PATH):
        os.remove(OUT_PATH)

    merchants, ledger = load_data()

    wb = Workbook()
    wb.remove(wb.active)

    write_merchants_sheet(wb, merchants)
    write_fee_tiers_sheet(wb)
    _, day_analysis_last_row = write_merchant_day_analysis_sheet(wb, merchants, ledger)
    write_transactions_view_sheet(wb, merchants, ledger, day_analysis_last_row)
    write_merchant_status_pivot_sheet(wb, merchants, ledger)

    wb.save(OUT_PATH)
    print(f"Saved: {OUT_PATH}")


if __name__ == "__main__":
    main()
